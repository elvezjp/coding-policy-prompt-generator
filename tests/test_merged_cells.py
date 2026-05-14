from pathlib import Path

from openpyxl import Workbook, load_workbook

from coding_policy_prompt_generator.excel_io import ColumnConfig, generate_prompts


COLUMNS = ColumnConfig(
    id_column="項番",
    summary_column="概要",
    description_column="説明",
    link_column="リンク",
)


def _write_header(ws) -> None:
    ws.cell(row=1, column=1, value="説明")
    ws.cell(row=2, column=1, value="説明2")
    headers = ["項番", "分類", "カテゴリ", "概要", "説明", "リンク"]
    for col, name in enumerate(headers, start=1):
        ws.cell(row=3, column=col, value=name)


def _run(input_path: Path, output_path: Path):
    return generate_prompts(
        input_path=input_path,
        output_path=output_path,
        index_sheet="コーディング規約一覧",
        header_row=3,
        columns=COLUMNS,
        sheet_prefix="PROMPT_",
        dry_run=False,
    )


def _user_prompt(wb, rule_id: str) -> str:
    return wb[f"PROMPT_{rule_id}"]["A2"].value


def test_classification_column_merged(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "コーディング規約一覧"
    _write_header(ws)

    for i, rid in enumerate(["N-001", "N-002", "N-003"], start=4):
        ws.cell(row=i, column=1, value=rid)
        ws.cell(row=i, column=3, value="クラス")
        ws.cell(row=i, column=4, value=f"概要{rid}")
    ws.cell(row=4, column=2, value="命名規則")
    ws.merge_cells(start_row=4, start_column=2, end_row=6, end_column=2)

    wb.save(input_path)
    _run(input_path, output_path)

    wb_out = load_workbook(output_path)
    for rid in ["N-001", "N-002", "N-003"]:
        prompt = _user_prompt(wb_out, rid)
        assert "| 分類 | 命名規則 |" in prompt, f"{rid}: {prompt}"


def test_category_column_merged(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "コーディング規約一覧"
    _write_header(ws)

    for i, rid in enumerate(["C-001", "C-002"], start=4):
        ws.cell(row=i, column=1, value=rid)
        ws.cell(row=i, column=2, value="命名規則")
        ws.cell(row=i, column=4, value=f"概要{rid}")
    ws.cell(row=4, column=3, value="クラス")
    ws.merge_cells(start_row=4, start_column=3, end_row=5, end_column=3)

    wb.save(input_path)
    _run(input_path, output_path)

    wb_out = load_workbook(output_path)
    for rid in ["C-001", "C-002"]:
        prompt = _user_prompt(wb_out, rid)
        assert "| カテゴリ | クラス |" in prompt, f"{rid}: {prompt}"


def test_description_column_merged(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "コーディング規約一覧"
    _write_header(ws)

    for i, rid in enumerate(["D-001", "D-002"], start=4):
        ws.cell(row=i, column=1, value=rid)
        ws.cell(row=i, column=2, value="命名規則")
        ws.cell(row=i, column=3, value="クラス")
        ws.cell(row=i, column=4, value=f"概要{rid}")
    ws.cell(row=4, column=5, value="共通の詳細ルール")
    ws.merge_cells(start_row=4, start_column=5, end_row=5, end_column=5)

    wb.save(input_path)
    _run(input_path, output_path)

    wb_out = load_workbook(output_path)
    for rid in ["D-001", "D-002"]:
        prompt = _user_prompt(wb_out, rid)
        assert "共通の詳細ルール" in prompt, f"{rid}: {prompt}"


def test_all_three_columns_merged(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "コーディング規約一覧"
    _write_header(ws)

    for i, rid in enumerate(["A-001", "A-002", "A-003"], start=4):
        ws.cell(row=i, column=1, value=rid)
        ws.cell(row=i, column=4, value=f"概要{rid}")
    ws.cell(row=4, column=2, value="命名規則")
    ws.cell(row=4, column=3, value="クラス")
    ws.cell(row=4, column=5, value="共通詳細")
    ws.merge_cells(start_row=4, start_column=2, end_row=6, end_column=2)
    ws.merge_cells(start_row=4, start_column=3, end_row=6, end_column=3)
    ws.merge_cells(start_row=4, start_column=5, end_row=6, end_column=5)

    wb.save(input_path)
    _run(input_path, output_path)

    wb_out = load_workbook(output_path)
    for rid in ["A-001", "A-002", "A-003"]:
        prompt = _user_prompt(wb_out, rid)
        assert "| 分類 | 命名規則 |" in prompt, f"{rid} classification: {prompt}"
        assert "| カテゴリ | クラス |" in prompt, f"{rid} category: {prompt}"
        assert "共通詳細" in prompt, f"{rid} description: {prompt}"


def test_top_left_of_merged_range_still_reads_value(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "コーディング規約一覧"
    _write_header(ws)

    ws.cell(row=4, column=1, value="T-001")
    ws.cell(row=4, column=2, value="命名規則")
    ws.cell(row=4, column=3, value="クラス")
    ws.cell(row=4, column=4, value="概要T-001")
    ws.merge_cells(start_row=4, start_column=2, end_row=5, end_column=2)

    wb.save(input_path)
    _run(input_path, output_path)

    wb_out = load_workbook(output_path)
    prompt = _user_prompt(wb_out, "T-001")
    assert "| 分類 | 命名規則 |" in prompt


def test_empty_cell_outside_merged_range_stays_unspecified(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "コーディング規約一覧"
    _write_header(ws)

    ws.cell(row=4, column=1, value="E-001")
    ws.cell(row=4, column=4, value="概要E-001")
    # 分類・カテゴリ列とも未入力かつ結合範囲外

    wb.save(input_path)
    _run(input_path, output_path)

    wb_out = load_workbook(output_path)
    prompt = _user_prompt(wb_out, "E-001")
    assert "| 分類 | （未指定） |" in prompt
    assert "| カテゴリ | （未指定） |" in prompt


def test_regression_no_merged_cells(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "コーディング規約一覧"
    _write_header(ws)

    for i, rid in enumerate(["R-001", "R-002"], start=4):
        ws.cell(row=i, column=1, value=rid)
        ws.cell(row=i, column=2, value="命名規則")
        ws.cell(row=i, column=3, value="クラス")
        ws.cell(row=i, column=4, value=f"概要{rid}")
        ws.cell(row=i, column=5, value=f"詳細{rid}")

    wb.save(input_path)
    _run(input_path, output_path)

    wb_out = load_workbook(output_path)
    for rid in ["R-001", "R-002"]:
        prompt = _user_prompt(wb_out, rid)
        assert "| 分類 | 命名規則 |" in prompt
        assert "| カテゴリ | クラス |" in prompt
        assert f"詳細{rid}" in prompt
